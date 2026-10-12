from __future__ import annotations

import re
import tempfile
from datetime import date
from pathlib import Path

import pypdfium2 as pdfium
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from build_delivery_pdf import CLASS_DIAGRAM_URL, DIAGRAMS, DOCS, ROOT, find_browser, render_diagram


OUTPUT = DOCS / "entrega-consolidada.docx"
MODEL_2_URL = CLASS_DIAGRAM_URL
CONTENTS = [
    "1. Modelos UML y C4",
    "2. Decisiones de arquitectura (ADR)",
    "3. Matriz de decisión del stack",
    "4. Plan de pruebas",
    "5. Prototipo funcionando",
    "6. Registro de deuda técnica",
    "7. Matriz de riesgos actualizada",
    "8. Tabla de trazabilidad",
    "9. Nota de uso de IA",
    "10. Documento consolidado y fuentes",
]
SECTION_SOURCES = {
    2: [
        DOCS / "adr/0001-stack-del-prototipo.md",
        DOCS / "adr/0002-logica-de-negocio-en-servicio.md",
        DOCS / "adr/0003-inventario-inicial-en-cero.md",
    ],
    3: [DOCS / "matriz-decision-stack.md"],
    4: [DOCS / "plan-pruebas.md"],
    5: [ROOT / "README.md"],
    6: [DOCS / "deuda-tecnica.md"],
    7: [DOCS / "matriz-riesgos.md"],
    8: [DOCS / "trazabilidad.md"],
    9: [DOCS / "nota-uso-ia.md"],
}


def clean_markdown(text: str) -> str:
    text = re.sub(r"!\[([^\]]*)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"\1", text)
    return text


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relation_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "176B55")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    run.append(properties)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_rich_text(paragraph, text: str) -> None:
    link_pattern = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
    position = 0
    for match in link_pattern.finditer(text):
        plain = clean_markdown(text[position:match.start()])
        if plain:
            paragraph.add_run(plain)
        add_hyperlink(paragraph, clean_markdown(match.group(1)), match.group(2))
        position = match.end()
    remainder = clean_markdown(text[position:])
    if remainder:
        paragraph.add_run(remainder)


def add_markdown_table(document: Document, lines: list[str]) -> None:
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if cells and all(re.fullmatch(r":?-{2,}:?", cell.replace(" ", "")) for cell in cells):
            continue
        rows.append(cells)
    if not rows:
        return

    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=1, cols=column_count)
    try:
        table.style = "Light Shading Accent 1"
    except KeyError:
        table.style = "Table Grid"
    table.autofit = True
    for row_index, row in enumerate(rows):
        cells = table.rows[0].cells if row_index == 0 else table.add_row().cells
        for column, text in enumerate(row):
            cell = cells[column]
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            paragraph = cell.paragraphs[0]
            add_rich_text(paragraph, text)
            for run in paragraph.runs:
                run.font.size = Pt(8)
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
        if row_index == 0:
            properties = table.rows[0]._tr.get_or_add_trPr()
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            properties.append(repeat)
    document.add_paragraph()


def add_markdown_file(document: Document, path: Path) -> None:
    lines = path.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line or line == "---":
            index += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            add_markdown_table(document, table_lines)
            continue

        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            if level > 1:
                document.add_heading(clean_markdown(heading.group(2)), level=min(level, 4))
        elif line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_rich_text(paragraph, line[2:])
        elif re.match(r"^\d+\.\s+", line):
            paragraph = document.add_paragraph(style="List Number")
            add_rich_text(paragraph, re.sub(r"^\d+\.\s+", "", line))
        elif line.startswith("> "):
            paragraph = document.add_paragraph(style="Quote")
            add_rich_text(paragraph, line[2:])
        else:
            paragraph = document.add_paragraph()
            add_rich_text(paragraph, line)
        index += 1


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in [begin, instruction, separate, value, end]:
        run._r.append(element)


def add_diagram(document: Document, caption: str, source_name: str, image_path: Path) -> None:
    document.add_heading(caption, level=2)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Inches(6.25))
    source_paragraph = document.add_paragraph()
    source_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    source_run = source_paragraph.add_run(f"Fuente editable: docs/modelos/{source_name}")
    source_run.italic = True
    source_run.font.size = Pt(8)
    document.add_page_break()


def build_docx() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(45, 58, 50)
    for heading_name, size, color in [
        ("Heading 1", 18, "174A35"),
        ("Heading 2", 14, "245B48"),
        ("Heading 3", 12, "40554C"),
    ]:
        style = document.styles[heading_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)

    header = section.header.paragraphs[0]
    header.text = "AgroStock | Entrega 2"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(102, 118, 110)
    add_page_number(section.footer.paragraphs[0])

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Inches(1.6)
    run = title.add_run("AgroStock")
    run.bold = True
    run.font.size = Pt(30)
    run.font.color.rgb = RGBColor(23, 74, 53)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Entrega 2\nIngeniería de Software").font.size = Pt(17)
    document.add_paragraph()
    team = document.add_paragraph()
    team.alignment = WD_ALIGN_PARAGRAPH.CENTER
    team.add_run("Equipo\n").bold = True
    team.add_run("Juan Felipe Cardona\nCamilo González\nDaniel Villa")
    dated = document.add_paragraph(date.today().isoformat())
    dated.alignment = WD_ALIGN_PARAGRAPH.CENTER
    repository = document.add_paragraph()
    repository.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_hyperlink(repository, "Repositorio: Camilogm1/Agrostock", "https://github.com/Camilogm1/Agrostock")

    document.add_page_break()
    document.add_heading("Índice", level=1)
    for item in CONTENTS:
        document.add_paragraph(item)

    browser = find_browser()
    with tempfile.TemporaryDirectory(prefix="agrostock-word-") as temporary:
        temp_dir = Path(temporary)
        diagram_images = {}
        for caption, filename in DIAGRAMS:
            diagram_pdf = render_diagram(DOCS / "modelos" / filename, temp_dir, browser)
            pdf = pdfium.PdfDocument(str(diagram_pdf))
            image_path = temp_dir / f"{Path(filename).stem}.png"
            try:
                image = pdf[0].render(scale=2.0).to_pil()
                image.save(image_path, format="PNG", optimize=True)
            finally:
                pdf.close()
            diagram_images[filename] = image_path

        for section_number, title_text in enumerate(CONTENTS, start=1):
            document.add_page_break()
            document.add_heading(title_text, level=1)
            if section_number == 1:
                document.add_heading("Modelo 2. Diagrama de clases", level=2)
                add_hyperlink(document.add_paragraph(), "Abrir Modelo 2 en Google Drive", MODEL_2_URL)
                document.add_paragraph(
                    "Para visualizar y editar el diagrama con facilidad, abre el enlace con draw.io (diagrams.net)."
                )
                document.add_paragraph("Las imágenes de los demás modelos son vistas para consulta; sus fuentes draw.io se conservan en docs/modelos/.")
                for caption, filename in DIAGRAMS:
                    if filename == "Agrostock-Diagrama de secuencia.drawio.svg":
                        title = "Modelo 3. Diagrama de secuencia"
                    elif filename == "Agrostock-Diagrama de componentes y despliegue.drawio.svg":
                        title = "Modelo 4. Componentes y despliegue"
                    elif filename == "Agrostock-C4 Nivel 1.drawio.svg":
                        title = "Modelo 5. Diagrama C4: nivel 1"
                    elif filename == "Agrostock- C4 Nivel 2.drawio.svg":
                        title = "Modelo 6. Diagrama C4: nivel 2"
                    else:
                        title = "Modelo 1. Diagrama de caso de uso"
                    add_diagram(document, title, filename, diagram_images[filename])
            elif section_number == 10:
                document.add_paragraph("Este archivo Word consolida los componentes de la entrega en formato editable. Para regenerarlo desde la raíz del repositorio:")
                document.add_paragraph("py -m pip install -r docs/requirements-word.txt\npy scripts/build_delivery_docx.py")
                document.add_paragraph("El Modelo 2 requiere acceso al enlace externo de Google Drive. Verifica que sus permisos permitan abrirlo.")
            else:
                for source in SECTION_SOURCES[section_number]:
                    add_markdown_file(document, source)

    document.core_properties.title = "AgroStock - Entrega 2"
    document.core_properties.author = "Equipo AgroStock"
    document.save(OUTPUT)
    print(f"Word generado: {OUTPUT.relative_to(ROOT)} ({OUTPUT.stat().st_size:,} bytes)")


if __name__ == "__main__":
    build_docx()